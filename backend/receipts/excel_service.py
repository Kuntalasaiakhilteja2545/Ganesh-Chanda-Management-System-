"""
Excel Export Service using openpyxl.
"""

import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


HEADER_FONT = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill(start_color='1A5276', end_color='1A5276', fill_type='solid')
HEADER_ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)
THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'),
)


def _style_header(ws, row, num_cols):
    """Apply header styling to a row."""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGNMENT
        cell.border = THIN_BORDER


def generate_donations_excel(donations, festival_id):
    """Generate Excel file for donations."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Donations'

    # Title
    from festivals.models import Festival
    try:
        festival = Festival.objects.get(id=festival_id)
        ws.merge_cells('A1:H1')
        ws['A1'] = f"Donations Report — {festival.name}"
        ws['A1'].font = Font(size=14, bold=True)
        ws['A1'].alignment = Alignment(horizontal='center')
    except Festival.DoesNotExist:
        pass

    # Headers
    headers = ['#', 'Receipt No', 'Donor Name', 'Amount', 'Payment',
               'Transaction ID', 'Date', 'Collected By']
    for col, header in enumerate(headers, 1):
        ws.cell(row=3, column=col, value=header)
    _style_header(ws, 3, len(headers))

    # Data
    total = 0
    for i, d in enumerate(donations, 1):
        row = i + 3
        receipt_no = d.receipt.receipt_number if hasattr(d, 'receipt') else ''
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=receipt_no)
        ws.cell(row=row, column=3, value=d.donor.name)
        ws.cell(row=row, column=4, value=d.amount)
        ws.cell(row=row, column=5, value=d.get_payment_method_display())
        ws.cell(row=row, column=6, value=d.transaction_id)
        ws.cell(row=row, column=7, value=d.donation_date.strftime('%d-%b-%Y'))
        ws.cell(row=row, column=8, value=d.collected_by.full_name or d.collected_by.username)
        total += d.amount

        # Apply borders
        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).border = THIN_BORDER

    # Total row
    total_row = donations.count() + 4
    ws.cell(row=total_row, column=3, value='TOTAL').font = Font(bold=True)
    ws.cell(row=total_row, column=4, value=total).font = Font(bold=True)

    # Column widths
    widths = [5, 18, 25, 12, 14, 20, 14, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def generate_expenses_excel(expenses, festival_id):
    """Generate Excel file for expenses."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Expenses'

    from festivals.models import Festival
    try:
        festival = Festival.objects.get(id=festival_id)
        ws.merge_cells('A1:G1')
        ws['A1'] = f"Expenses Report — {festival.name}"
        ws['A1'].font = Font(size=14, bold=True)
        ws['A1'].alignment = Alignment(horizontal='center')
    except Festival.DoesNotExist:
        pass

    headers = ['#', 'Category', 'Description', 'Amount', 'Payment', 'Date', 'Paid By']
    for col, header in enumerate(headers, 1):
        ws.cell(row=3, column=col, value=header)
    _style_header(ws, 3, len(headers))

    total = 0
    for i, e in enumerate(expenses, 1):
        row = i + 3
        ws.cell(row=row, column=1, value=i)
        ws.cell(row=row, column=2, value=e.category.name)
        ws.cell(row=row, column=3, value=e.description)
        ws.cell(row=row, column=4, value=e.amount)
        ws.cell(row=row, column=5, value=e.get_payment_method_display())
        ws.cell(row=row, column=6, value=e.expense_date.strftime('%d-%b-%Y'))
        ws.cell(row=row, column=7, value=e.paid_by.full_name or e.paid_by.username)
        total += e.amount

        for col in range(1, len(headers) + 1):
            ws.cell(row=row, column=col).border = THIN_BORDER

    total_row = expenses.count() + 4
    ws.cell(row=total_row, column=3, value='TOTAL').font = Font(bold=True)
    ws.cell(row=total_row, column=4, value=total).font = Font(bold=True)

    widths = [5, 20, 30, 12, 14, 14, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
